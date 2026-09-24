import io
import uuid

import docx
import pypdf
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.docs.chunker import chunk_text
from app.modules.docs.embedder import Embedder
from app.modules.docs.models import Chunk, Document, DocumentStatus
from app.modules.search.base import SearchResult


class DocumentNotFoundError(Exception):
    pass


async def create_document(
    db: AsyncSession, user_id: uuid.UUID, filename: str, content_type: str
) -> Document:
    document = Document(
        user_id=user_id, filename=filename, content_type=content_type, status=DocumentStatus.PENDING
    )
    db.add(document)
    await db.flush()
    return document


def extract_text(content_type: str, raw_bytes: bytes) -> list[tuple[str, int | None]]:
    """Returns (text_segment, page_number) pairs. page_number is 1-indexed, None for non-paged formats."""
    if content_type == "application/pdf":
        reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
        return [(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(io.BytesIO(raw_bytes))
        return [("\n".join(p.text for p in doc.paragraphs), None)]
    return [(raw_bytes.decode("utf-8"), None)]


async def _chunk_embed_and_store(
    db: AsyncSession, document_id: uuid.UUID, segments: list[tuple[str, int | None]], embedder: Embedder
) -> None:
    """segments: list of (text, page_number). Shared by document upload ingestion and web search ingestion."""
    pieces = [(piece, page_number) for text, page_number in segments for piece in chunk_text(text)]

    embeddings = await embedder.embed([text for text, _ in pieces]) if pieces else []

    for index, ((text, page_number), embedding) in enumerate(zip(pieces, embeddings, strict=True)):
        db.add(
            Chunk(
                document_id=document_id,
                chunk_index=index,
                content=text,
                page_number=page_number,
                embedding=embedding,
            )
        )


async def ingest_document(
    db: AsyncSession,
    document_id: uuid.UUID,
    raw_bytes: bytes,
    content_type: str,
    embedder: Embedder,
) -> None:
    """Owns its transaction (worker calls this with its own session, no HTTP request around it).

    Whole body is wrapped so a Document is never left stuck in PROCESSING: any
    exception here downgrades it to FAILED with the error recorded instead.
    """
    document = await db.get(Document, document_id)
    if document is None:
        return

    try:
        document.status = DocumentStatus.PROCESSING
        await db.flush()

        segments = extract_text(content_type, raw_bytes)
        await _chunk_embed_and_store(db, document_id, segments, embedder)

        document.status = DocumentStatus.READY
        await db.commit()
    except Exception as exc:
        await db.rollback()
        document = await db.get(Document, document_id)
        document.status = DocumentStatus.FAILED
        document.error_message = str(exc)
        await db.commit()


async def ingest_search_results(
    db: AsyncSession, user_id: uuid.UUID, results: list[SearchResult], embedder: Embedder
) -> list[Document]:
    """One Document per SearchResult (not one per search call): Document.source_url is a single
    column, so a shared Document can't hold >1 URL. Per-result Document keeps each Chunk traceable
    to its exact source URL for citations, with no schema change needed."""
    documents = [
        Document(
            user_id=user_id,
            filename=result.title or result.url,
            content_type="text/html",
            status=DocumentStatus.READY,
            source_url=result.url,
        )
        for result in results
    ]
    db.add_all(documents)
    await db.flush()

    for document, result in zip(documents, results, strict=True):
        await _chunk_embed_and_store(db, document.id, [(result.content, None)], embedder)

    return documents


async def list_documents(db: AsyncSession, user_id: uuid.UUID) -> list[Document]:
    """User-uploaded documents only -- web-search-ingested pages (source_url
    set, see ingest_search_results) stay in the DB for retrieval but were
    never something the user chose to manage, so they don't belong in a doc
    list/picker UI."""
    result = await db.scalars(
        select(Document)
        .where(Document.user_id == user_id, Document.source_url.is_(None))
        .order_by(Document.created_at.desc())
    )
    return list(result.all())


async def get_document(db: AsyncSession, user_id: uuid.UUID, document_id: uuid.UUID) -> Document:
    document = await db.get(Document, document_id)
    if document is None or document.user_id != user_id:
        raise DocumentNotFoundError(document_id)
    return document


async def delete_document(db: AsyncSession, user_id: uuid.UUID, document_id: uuid.UUID) -> None:
    document = await get_document(db, user_id, document_id)
    await db.delete(document)
    await db.flush()


async def hybrid_search(
    db: AsyncSession,
    user_id: uuid.UUID,
    query: str,
    embedder: Embedder,
    top_k: int = 5,
    document_ids: list[uuid.UUID] | None = None,
    exclude_web_origin: bool = False,
) -> list[Chunk]:
    """Reciprocal Rank Fusion over vector similarity + Postgres full-text search.

    document_ids optionally scopes the search to a specific set of documents (e.g. the
    ones just ingested from a single web search call), instead of every READY document
    the user has.

    exclude_web_origin drops documents that came from a web search (source_url set) --
    for use_docs, which means "ground on my uploaded files", not on some unrelated web
    search from an earlier turn.
    """
    query_vector = (await embedder.embed([query]))[0]

    def _scoped(stmt):
        stmt = stmt.where(Document.user_id == user_id, Document.status == DocumentStatus.READY)
        if document_ids is not None:
            stmt = stmt.where(Document.id.in_(document_ids))
        if exclude_web_origin:
            stmt = stmt.where(Document.source_url.is_(None))
        return stmt

    vector_hits = (
        await db.scalars(
            _scoped(select(Chunk).join(Document))
            .order_by(Chunk.embedding.cosine_distance(query_vector))
            .limit(20)
        )
    ).all()

    fts_hits = (
        await db.scalars(
            _scoped(select(Chunk).join(Document))
            .where(text("to_tsvector('english', chunks.content) @@ plainto_tsquery('english', :q)"))
            .params(q=query)
            .limit(20)
        )
    ).all()

    # Reciprocal Rank Fusion, k=60 (standard constant)
    scores: dict[uuid.UUID, float] = {}
    chunks_by_id: dict[uuid.UUID, Chunk] = {}
    for rank_list in (vector_hits, fts_hits):
        for rank, chunk in enumerate(rank_list):
            scores[chunk.id] = scores.get(chunk.id, 0.0) + 1.0 / (60 + rank)
            chunks_by_id[chunk.id] = chunk

    ranked_ids = sorted(scores, key=lambda cid: scores[cid], reverse=True)
    return [chunks_by_id[cid] for cid in ranked_ids[:top_k]]
